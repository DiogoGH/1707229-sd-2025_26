from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import Response
from docx import Document
import pdfplumber
import io
import os


app = FastAPI(
    title="Web Service de Extração de Texto",
    description="Serviço REST para extração de texto de documentos PDF e DOCX.",
    version="1.0.0"
)


def extrair_texto_pdf(file_bytes: bytes) -> str:
    """
    Função responsável por extrair texto de um ficheiro PDF.
    Recebe o ficheiro em bytes e devolve o texto extraído.
    """

    texto_extraido = ""

    # Abre o PDF a partir dos bytes recebidos no upload
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:

        # Percorre todas as páginas do PDF
        for numero_pagina, pagina in enumerate(pdf.pages, start=1):

            # Extrai o texto da página atual
            texto_pagina = pagina.extract_text()

            # Se existir texto na página, adiciona ao resultado final
            if texto_pagina:
                texto_extraido += f"\n--- Página {numero_pagina} ---\n"
                texto_extraido += texto_pagina + "\n"

    return texto_extraido.strip()


def extrair_texto_docx(file_bytes: bytes) -> str:
    """
    Função responsável por extrair texto de um ficheiro DOCX.
    Lê os parágrafos e as tabelas existentes no documento.
    """

    texto_extraido = ""

    # Abre o documento DOCX a partir dos bytes recebidos no upload
    documento = Document(io.BytesIO(file_bytes))

    # Extrai o texto dos parágrafos
    for paragrafo in documento.paragraphs:
        if paragrafo.text.strip():
            texto_extraido += paragrafo.text + "\n"

    # Extrai o texto das tabelas, caso existam
    for tabela in documento.tables:
        for linha in tabela.rows:
            valores_linha = []

            for celula in linha.cells:
                valores_linha.append(celula.text.strip())

            texto_extraido += " | ".join(valores_linha) + "\n"

    return texto_extraido.strip()


@app.post("/extract/text")
async def extract_text(file: UploadFile = File(...)):
    """
    Endpoint responsável por receber um ficheiro PDF ou DOCX
    e devolver um ficheiro TXT com o texto extraído.
    """

    # Verifica se foi enviado algum ficheiro
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="Nenhum ficheiro foi enviado."
        )

    # Guarda o nome original do ficheiro
    nome_ficheiro = file.filename

    # Obtém a extensão do ficheiro
    extensao = os.path.splitext(nome_ficheiro)[1].lower()

    # Valida se o ficheiro é PDF ou DOCX
    if extensao not in [".pdf", ".docx"]:
        raise HTTPException(
            status_code=400,
            detail="Formato inválido. Apenas são permitidos ficheiros PDF ou DOCX."
        )

    # Lê o conteúdo do ficheiro enviado
    file_bytes = await file.read()

    # Verifica se o ficheiro está vazio
    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="O ficheiro enviado está vazio."
        )

    try:
        # Extrai texto conforme o tipo de ficheiro recebido
        if extensao == ".pdf":
            texto = extrair_texto_pdf(file_bytes)

        elif extensao == ".docx":
            texto = extrair_texto_docx(file_bytes)

    except Exception as erro:
        raise HTTPException(
            status_code=500,
            detail=f"Erro ao processar o ficheiro: {str(erro)}"
        )

    # Verifica se foi possível extrair algum texto
    if not texto.strip():
        raise HTTPException(
            status_code=400,
            detail="Não foi possível extrair texto do documento. O ficheiro pode estar vazio ou conter apenas imagens."
        )

    # Cria o nome do ficheiro TXT de resposta
    nome_base = os.path.splitext(nome_ficheiro)[0]
    nome_txt = f"{nome_base}.txt"

    # Devolve o ficheiro TXT ao cliente
    return Response(
        content=texto,
        media_type="text/plain; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{nome_txt}"'
        }
    )

"py -m uvicorn main:app --reload"